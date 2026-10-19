"""Security Cloud findings: adversarial inputs and ordinary authenticated flows."""
import io
import os
import ssl
import threading
import zipfile
from concurrent.futures import Future
from unittest.mock import MagicMock

import openpyxl
import pytest
from werkzeug.datastructures import FileStorage

from database import db
from models import Team, Tournament
from models.competitor import ProCompetitor
from models.school_captain import SchoolCaptain
from models.user import User
from tests.db_test_utils import create_test_app, drop_test_db


@pytest.fixture
def security_app():
    app, path = create_test_app()
    with app.app_context():
        tournament = Tournament(name='Security tournament', year=2027)
        admin = User(username='security-admin', role='admin')
        admin.set_password('long-test-password')
        registrar = User(username='security-registrar', role='registrar')
        registrar.set_password('long-test-password')
        db.session.add_all([tournament, admin, registrar])
        db.session.flush()
        team = Team(tournament_id=tournament.id, school_name='Test School', school_abbreviation='TS', team_code='TS-A')
        competitor = ProCompetitor(tournament_id=tournament.id, name='Test Competitor', gender='M')
        db.session.add_all([team, competitor])
        db.session.commit()
        app.config.update(TID=tournament.id, CID=competitor.id, TEAM_ID=team.id, ADMIN_ID=admin.id, REG_ID=registrar.id)
    yield app
    with app.app_context():
        db.session.remove()
        db.engine.dispose()
    drop_test_db(path)


def _login_as(app, user_id):
    client = app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(user_id)
        session['_fresh'] = True
    return client


def test_empty_database_http_bootstrap_is_closed_and_cli_works():
    app, path = create_test_app()
    try:
        client = app.test_client()
        assert client.get('/auth/bootstrap').status_code == 404
        assert client.post('/auth/bootstrap', data={'username': 'attacker', 'password': 'password123'}).status_code == 404
        runner = app.test_cli_runner()
        response = runner.invoke(args=['auth', 'create-admin'], input='trusted-admin\nlong-password-123\nlong-password-123\n')
        assert response.exit_code == 0, response.output
        assert runner.invoke(args=['auth', 'create-admin']).exit_code != 0
        with app.app_context():
            assert User.query.one().username == 'trusted-admin'
    finally:
        with app.app_context():
            db.session.remove()
            db.engine.dispose()
        drop_test_db(path)


def test_disabled_existing_session_is_rejected(security_app):
    client = _login_as(security_app, security_app.config['ADMIN_ID'])
    with security_app.app_context():
        user = db.session.get(User, security_app.config['ADMIN_ID'])
        user.is_active_user = False
        db.session.commit()
        # Pinned Flask-Login already binds authenticated to active; keep the
        # loader check as defense in depth against future mixin changes.
        assert user.is_authenticated is False
    assert client.get('/auth/users').status_code == 302
    assert client.post('/registration/1/college/team/1/override-validation').status_code == 302


def test_unclaimed_portals_require_admin_provisioning(security_app):
    app = security_app
    public = app.test_client()
    data = {'tournament_id': app.config['TID'], 'competitor_type': 'pro', 'competitor_id': app.config['CID'], 'pin': '12345678', 'confirm_pin': '12345678'}
    assert public.post('/portal/competitor/claim', data=data).status_code == 403
    school = {'tournament_id': app.config['TID'], 'school_name': 'Test School', 'pin': '12345678', 'confirm_pin': '12345678'}
    assert public.post('/portal/school/claim', data=school).status_code == 403
    assert public.post('/auth/users/portal-pin', data=data).status_code == 302
    with app.app_context():
        assert not db.session.get(ProCompetitor, app.config['CID']).portal_pin_hash
        assert SchoolCaptain.query.count() == 0


@pytest.mark.parametrize('kind', ['pro', 'school'])
def test_admin_provisioned_pin_unlocks_and_reset_revokes(security_app, monkeypatch, kind):
    app = security_app
    admin = _login_as(app, app.config['ADMIN_ID'])
    monkeypatch.setattr('secrets.randbelow', lambda n: 12345678)
    data = {'tournament_id': app.config['TID'], 'competitor_type': kind, 'competitor_id': app.config['CID'], 'school_name': 'Test School', 'identity_verified': 'yes'}
    response = admin.post('/auth/users/portal-pin', data={**data, 'profile': 'school:Test School' if kind == 'school' else f"pro:{app.config['CID']}"})
    assert response.status_code == 200
    assert b'12345678' in response.data and response.headers['Cache-Control'] == 'no-store'
    public = app.test_client()
    endpoint = '/portal/school/claim' if kind == 'school' else '/portal/competitor/claim'
    response = public.post(endpoint, data={**data, 'pin': '12345678'})
    assert response.status_code == 302
    assert public.get(response.location).status_code == 200
    old_location = response.location
    monkeypatch.setattr('secrets.randbelow', lambda n: 87654321)
    assert admin.post('/auth/users/portal-pin', data=data).status_code == 200
    assert public.get(old_location).status_code == 302


def test_rate_limits_survive_factory_and_rotating_addresses(security_app):
    app = security_app
    app.config['RATELIMIT_ENABLED'] = True
    from routes.api import _limiter
    with app.app_context():
        _limiter.reset()
    client = app.test_client()
    for i in range(5):
        assert client.post('/auth/login', data={'username': 'security-admin', 'password': 'wrong'}, environ_overrides={'REMOTE_ADDR': f'192.0.2.{i}'}).status_code == 200
    assert client.post('/auth/login', data={'username': 'security-admin', 'password': 'wrong'}, environ_overrides={'REMOTE_ADDR': '192.0.2.99'}).status_code == 429


def test_registrar_cannot_override_but_admin_reason_is_audited(security_app):
    app = security_app
    base = f"/registration/{app.config['TID']}/college/team/{app.config['TEAM_ID']}"
    registrar = _login_as(app, app.config['REG_ID'])
    for action in ['override-validation', 'remove-override']:
        assert registrar.post(base + '/' + action, data={'reason': 'test'}).status_code == 403
    admin = _login_as(app, app.config['ADMIN_ID'])
    assert admin.post(base + '/override-validation', data={'reason': 'Verified exception'}).status_code == 302
    with app.app_context():
        from models.audit_log import AuditLog
        record = AuditLog.query.filter_by(action='team_validation_override').one()
        assert record.actor_user_id == app.config['ADMIN_ID']
        assert 'Verified exception' in record.details_json


def test_stream_rejects_immediately_and_polling_remains(security_app):
    client = security_app.test_client()
    assert client.get('/api/public/tournaments/999999/standings-stream').status_code == 404
    tid = security_app.config['TID']
    assert client.get(f'/api/public/tournaments/{tid}/standings-stream').status_code == 410
    assert client.get(f'/api/public/tournaments/{tid}/standings-poll').status_code == 200
    assert client.get('/health').status_code == 200


@pytest.mark.parametrize('referrer', ['https://evil.example/', '//evil.example/', '/\\evil.example/', 'https://['])
def test_csrf_failure_cannot_redirect_off_origin(security_app, referrer):
    security_app.config.update(WTF_CSRF_ENABLED=True, WTF_CSRF_CHECK_DEFAULT=True)
    response = security_app.test_client().post('/auth/login', headers={'Referer': referrer})
    assert response.status_code == 302
    assert response.location == '/auth/login'


@pytest.mark.parametrize('member,content', [
    ('xl/worksheets/sheet1.xml', b'<worksheet><dimension ref="A1:XFD1048576"/></worksheet>'),
    ('xl/sharedStrings.xml', b'A' * 1000000),
    ('xl/workbook.xml', '<!DOCTYPE x [<!ENTITY z "expanded">]><x>&z;</x>'.encode('utf-16')),
])
def test_workbook_expansion_and_dimensions_are_rejected(member, content):
    from services.upload_security import validate_excel_upload
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, content)
    stream.seek(0)
    result = validate_excel_upload(FileStorage(stream=stream, filename='entries.xlsx'), {'xlsx'})
    assert not result.ok
    assert stream.tell() == 0


def test_literal_xlsx_cells_preserve_malicious_text(tmp_path):
    from services.spreadsheet_security import literal_workbook
    workbook = openpyxl.Workbook()
    payloads = ['=HYPERLINK("https://evil.example")', '+SUM(1)', '-1+1', '@SUM(1)', '\t=1', '\r=1']
    for value in payloads:
        workbook.active.append([value])
    literal_workbook(workbook)
    path = tmp_path / 'literal.xlsx'
    workbook.save(path)
    loaded = openpyxl.load_workbook(path)
    assert [row[0].value for row in loaded.active] == [value.replace('\r', '\n') for value in payloads]
    assert all(row[0].data_type == 's' for row in loaded.active)


def test_smtp_tls_verifies_peer_and_fails_before_login(monkeypatch):
    from services import email_delivery
    for key, value in {'SMTP_HOST': 'smtp.example', 'SMTP_USER': 'sender', 'SMTP_PASSWORD': 'secret'}.items():
        monkeypatch.setenv(key, value)
    server = MagicMock()
    server.__enter__.return_value = server
    server.starttls.side_effect = ssl.SSLCertVerificationError('wrong hostname')
    monkeypatch.setattr(email_delivery.smtplib, 'SMTP', lambda *args, **kwargs: server)
    result = email_delivery.send_document(['a@example.com'], 'test', 'body', b'data', 'test.pdf')
    assert result.status == 'failed'
    context = server.starttls.call_args.kwargs['context']
    assert context.verify_mode == ssl.CERT_REQUIRED and context.check_hostname
    server.login.assert_not_called()
    server.send_message.assert_not_called()


def test_backup_preserves_tls_without_credentials_in_argv():
    from services.backup import _pg_dump_args_and_env
    args, env = _pg_dump_args_and_env('postgresql://operator:secret@db.example/data?sslmode=verify-full&sslrootcert=%2Ftrusted%2Froot.pem&sslcert=%2Fclient.pem&connect_timeout=5')
    assert 'secret' not in ' '.join(args)
    assert env['PGSSLMODE'] == 'verify-full'
    assert env['PGSSLROOTCERT'] == '/trusted/root.pem'
    assert env['PGSSLCERT'] == '/client.pem'
    assert env['PGCONNECT_TIMEOUT'] == '5'


def test_job_queue_has_global_and_user_bounds(monkeypatch):
    from services import background_jobs as jobs
    monkeypatch.setattr(jobs, '_app', None)
    monkeypatch.setattr(jobs, '_jobs', {})
    executor = MagicMock()
    executor.submit.side_effect = lambda *args, **kwargs: Future()
    monkeypatch.setattr(jobs, '_executor', executor)
    for _ in range(jobs.MAX_PENDING_PER_USER):
        jobs.submit('test', lambda: None, metadata={'user_id': 1})
    with pytest.raises(jobs.JobQueueFull):
        jobs.submit('test', lambda: None, metadata={'user_id': 1})
    for i in range(jobs.MAX_PENDING_JOBS - jobs.MAX_PENDING_PER_USER):
        jobs.submit('test', lambda: None, metadata={'user_id': i + 2})
    with pytest.raises(jobs.JobQueueFull):
        jobs.submit('test', lambda: None)
    assert len(jobs._jobs) == jobs.MAX_PENDING_JOBS


@pytest.mark.parametrize('kind', ['college', 'pro'])
@pytest.mark.parametrize('fail', [False, True])
def test_original_upload_is_removed_on_every_parse_outcome(security_app, tmp_path, monkeypatch, kind, fail):
    app = security_app
    app.config['UPLOAD_FOLDER'] = str(tmp_path)
    client = _login_as(app, app.config['ADMIN_ID'])
    workbook = openpyxl.Workbook()
    workbook.active.append(['Name', 'Gender'])
    workbook.active.append(['Example', 'M'])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    def parse(*args, **kwargs):
        assert os.path.isfile(args[0])
        if fail:
            raise ValueError('parse failure')
        return {'teams': 0, 'competitors': 0} if kind == 'college' else []
    if kind == 'college':
        monkeypatch.setattr('services.excel_io.process_college_entry_form', parse)
        endpoint = f"/registration/{app.config['TID']}/college/upload"
    else:
        monkeypatch.setattr('services.pro_entry_importer.parse_pro_entries', parse)
        endpoint = f"/import/{app.config['TID']}/pro-entries"
    assert client.post(endpoint, data={'file': (stream, 'entries.xlsx')}).status_code == 302
    assert list(tmp_path.glob('*.xlsx')) == []


def test_all_xlsx_writers_keep_formula_text_literal(security_app, tmp_path, monkeypatch):
    from types import SimpleNamespace

    from services import excel_io, handicap_export, video_judge_export
    payload = '=1+1'
    with security_app.app_context():
        tournament = db.session.get(Tournament, security_app.config['TID'])
        team = db.session.get(Team, security_app.config['TEAM_ID'])
        team.school_name = payload
        team.team_code = payload
        db.session.commit()
        results = tmp_path / 'results.xlsx'
        excel_io.export_results_to_excel(tournament, str(results))
        monkeypatch.setattr(handicap_export, 'build_chopping_rows', lambda _: [{'Name': payload, 'Result': 7.5}])
        chopping = tmp_path / 'chopping.xlsx'
        handicap_export.export_chopping_results_to_excel(tournament, str(chopping))
    event = SimpleNamespace(event_type='pro', scoring_type='time', requires_dual_runs=False, requires_triple_runs=False)
    video = tmp_path / 'video.xlsx'
    video_judge_export.write_workbook({'Test': {'event': event, 'rows': [{
        'heat_number': 1, 'run_number': 1, 'stand_display': '1', 'competitor_name': payload,
        'team_code': None,
    }]}}, str(video))
    for path in [results, chopping, video]:
        workbook = openpyxl.load_workbook(path)
        cells = [cell for sheet in workbook for row in sheet for cell in row]
        assert any(cell.value == payload for cell in cells)
        assert all(cell.data_type != 'f' for cell in cells)


def test_email_queue_rejection_never_sends_and_attachment_size_is_bounded(security_app, monkeypatch):
    from services import background_jobs, email_delivery
    worker = MagicMock()
    monkeypatch.setattr(email_delivery, '_worker_send', worker)
    monkeypatch.setattr(background_jobs, 'submit', MagicMock(side_effect=background_jobs.JobQueueFull()))
    with security_app.app_context():
        args = dict(tournament_id=security_app.config['TID'], doc_key='test', entity_id=None,
                    recipients=['a@example.com'], subject='test', body='test', attachment_bytes=b'x', attachment_name='test.pdf')
        with pytest.raises(background_jobs.JobQueueFull):
            email_delivery.queue_document_email(**args)
        worker.assert_not_called()
        from werkzeug.exceptions import BadRequest
        args['attachment_bytes'] = b'x' * (email_delivery.MAX_ATTACHMENT_BYTES + 1)
        with pytest.raises(BadRequest):
            email_delivery.queue_document_email(**args)
        assert background_jobs.submit.call_count == 1


def test_export_and_pin_limits_return_429(security_app, monkeypatch):
    app = security_app
    app.config['RATELIMIT_ENABLED'] = True
    from routes.api import _limiter
    with app.app_context():
        _limiter.reset()
        competitor = db.session.get(ProCompetitor, app.config['CID'])
        competitor.set_portal_pin('12345678')
        db.session.commit()
    public = app.test_client()
    data = {'tournament_id': app.config['TID'], 'competitor_type': 'pro', 'competitor_id': app.config['CID'], 'pin': '0000'}
    for i in range(5):
        assert public.post('/portal/competitor/claim', data=data, environ_overrides={'REMOTE_ADDR': f'198.51.100.{i}'}).status_code == 302
    assert public.post('/portal/competitor/claim', data=data, environ_overrides={'REMOTE_ADDR': '198.51.100.99'}).status_code == 429
    admin = _login_as(app, app.config['ADMIN_ID'])
    monkeypatch.setattr('routes.reporting.build_chopping_json_payload', lambda _: {})
    url = f"/reporting/{app.config['TID']}/export-chopping?format=json"
    assert [admin.get(url).status_code for _ in range(4)] == [200, 200, 200, 429]


def test_terminal_jobs_expire_from_memory_and_database(security_app, monkeypatch):
    from datetime import timedelta

    from models.background_job import BackgroundJob
    from services import background_jobs as jobs
    from services.time_utils import utc_now_naive
    old = utc_now_naive() - timedelta(days=2)
    with security_app.app_context():
        db.session.add(BackgroundJob(id='expired', label='old', status='completed', submitted_at=old))
        db.session.commit()
    monkeypatch.setattr(jobs, '_app', security_app)
    monkeypatch.setattr(jobs, '_jobs', {'expired': {
        'id': 'expired', 'status': 'completed', 'submitted_at': old, 'terminal_persisted': True,
    }})
    jobs._prune_terminal_jobs()
    assert jobs._jobs == {}
    with security_app.app_context():
        assert db.session.get(BackgroundJob, 'expired') is None


def test_upload_save_failure_removes_partial_file(tmp_path):
    from services.upload_security import save_upload
    file = MagicMock()
    def fail(destination):
        destination.write(b'partial upload')
        raise OSError('disk failure')
    file.save.side_effect = fail
    with pytest.raises(OSError):
        save_upload(file, str(tmp_path), 'temporary.xlsx')
    assert not (tmp_path / 'temporary.xlsx').exists()
